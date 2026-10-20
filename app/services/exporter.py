"""Excel export.

Reproduces the layout of `Database Format.xlsx`: the seven bilingual columns in
their original order, followed by operational columns so the export carries the
complete database ("The Database should be able to be exported into an excel
with all columns by administrator roles").

Accuracy and completeness are checked, not assumed:

* the rows come from the same query as the Standards list
  (`standards_query`), so an export taken with the list's filters holds
  exactly the rows the list shows, in the same order;
* scraped text is written as text - a title beginning with "=" never becomes
  a formula, and control characters that Excel rejects are stripped instead
  of aborting the whole export;
* `verify_export` re-reads the saved workbook and compares it, row by row and
  field by field, with a fresh query; the result is written to the
  "Export Check" sheet and returned to the caller.
"""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.config import settings
from app.models import ChangeRecord, Source, Standard
from app.services.standards_query import standards_query

#: (header, model attribute, column width) — the first seven match the
#: bilingual headers of Database Format.xlsx exactly.
COLUMNS: list[tuple[str, str, int]] = [
    ("标准号+年份\nStandard ID + Version", "standard_id", 34),
    ("归口单位代码\nTC / SC Code", "tc_code", 20),
    ("标准中文名称\nStandard Chinese Title", "title_zh", 30),
    ("标准英文名称\nStandard English Title", "title_en", 70),
    ("发布日期\nPublish Date", "publish_date", 16),
    ("实施日期\nImplement Date", "implement_date", 16),
    ("标准状态\nStatus", "status", 14),
    ("标准组织\nStandards Body", "_body", 16),
    ("出版物类型\nPublication Type", "publication_type", 22),
    ("标准系列\nStandard Family", "family", 24),
    ("替代\nSupersedes", "supersedes", 30),
    ("被替代\nSuperseded By", "superseded_by", 30),
    ("最近变更\nLast Change", "last_change_type", 14),
    ("最近变更运行\nLast Change Run", "last_change_run_id", 12),
    ("首次运行\nFirst Seen Run", "first_seen_run_id", 12),
    ("最后运行\nLast Seen Run", "last_seen_run_id", 12),
    ("废止运行\nWithdrawn In Run", "withdrawn_run_id", 12),
    ("数据来源\nSource URL", "source_url", 46),
    ("首次登记\nFirst Captured", "_created", 20),
    ("最后更新\nLast Updated", "_updated", 20),
]

#: Fields compared cell-for-cell by `verify_export`.
VERIFIED_FIELDS = ("standard_id", "tc_code", "title_zh", "title_en", "publish_date",
                   "implement_date", "status")

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
CHECK_SHEET = "Export Check"


def safe_cell(value):
    """Text exactly as stored, minus characters Excel cannot hold."""
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def _append(sheet, values) -> None:
    sheet.append([safe_cell(v) for v in values])
    # openpyxl turns any string starting with "=" into a formula; scraped text
    # must stay text (this also blocks formula injection from a source site).
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str) and cell.value.startswith("="):
            cell.data_type = "s"


def export_workbook(
    db: Session,
    *,
    q: str | None = None,
    status: str | None = None,
    body: str | None = None,
    tc_code: str | None = None,
    publication_type: str | None = None,
    change_type: str | None = None,
    include_change_log: bool = True,
) -> tuple[Path, dict]:
    """Write the workbook, verify it against the database, return (path, report)."""
    filters = {
        "q": q,
        "status": status,
        "body": body,
        "tc_code": tc_code,
        "publication_type": publication_type,
        "change_type": change_type,
    }
    rows = list(db.execute(standards_query(**filters)))

    workbook = Workbook()
    _write_standards(workbook.active, rows)
    _write_sources(workbook.create_sheet("Sources"), db)
    if include_change_log:
        _write_change_log(workbook.create_sheet("Change Log"), db)

    settings.export_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
    path = settings.export_dir / f"standards_database_{stamp}.xlsx"
    workbook.save(path)

    report = verify_export(db, path, **filters)
    _write_check_sheet(path, report)
    return path, report


def verify_export(db: Session, path: Path, **filters) -> dict:
    """Re-read a saved export and compare it with the database, row by row."""
    expected = [standard for standard, _ in db.execute(standards_query(**filters))]
    workbook = load_workbook(path, read_only=True)
    try:
        rows = workbook["Standards"].iter_rows(values_only=True)
        header = list(next(rows, ()) or ())
        exported = [row for row in rows if any(v not in (None, "") for v in row)]
    finally:
        workbook.close()

    index = {attr: header.index(name) for name, attr, _ in COLUMNS if name in header}
    missing_columns = [name for name, _, _ in COLUMNS if name not in header]

    def cell(row, attr):
        position = index.get(attr)
        if position is None or position >= len(row) or row[position] in (None, ""):
            return None
        return str(row[position])

    mismatches: list[dict] = []
    for position, standard in enumerate(expected[: len(exported)]):
        row = exported[position]
        for attr in VERIFIED_FIELDS:
            if attr not in index:
                continue
            wanted = safe_cell(getattr(standard, attr)) or None
            got = cell(row, attr)
            if wanted != got:
                mismatches.append(
                    {"row": position + 2, "standard_id": standard.standard_id,
                     "field": attr, "database": wanted, "export": got}
                )

    expected_ids = [s.standard_id for s in expected]
    exported_ids = [cell(row, "standard_id") for row in exported]
    missing = sorted(set(expected_ids) - set(exported_ids))
    extra = sorted(str(i) for i in set(exported_ids) - set(expected_ids))

    by_status_db: dict[str, int] = {}
    for standard in expected:
        by_status_db[standard.status] = by_status_db.get(standard.status, 0) + 1
    by_status_export: dict[str, int] = {}
    for row in exported:
        key = cell(row, "status") or ""
        by_status_export[key] = by_status_export.get(key, 0) + 1

    ok = (
        len(exported) == len(expected)
        and not missing
        and not extra
        and not mismatches
        and not missing_columns
    )
    return {
        "ok": ok,
        "file": path.name,
        "checked_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "filters": {k: v for k, v in filters.items() if v},
        "database_rows": len(expected),
        "database_total": db.scalar(select(func.count()).select_from(Standard)) or 0,
        "exported_rows": len(exported),
        "by_status": {"database": by_status_db, "export": by_status_export},
        "missing_from_export": missing[:100],
        "unexpected_in_export": extra[:100],
        "missing_columns": missing_columns,
        "field_mismatches": mismatches[:100],
        "field_mismatch_count": len(mismatches),
    }


def _write_check_sheet(path: Path, report: dict) -> None:
    workbook = load_workbook(path)
    sheet = workbook.create_sheet(CHECK_SHEET)
    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 70
    filters = ", ".join(f"{k}={v}" for k, v in report["filters"].items()) or "none (full database)"
    lines = [
        ("Result", "PASS - export matches the standards list" if report["ok"] else "FAIL"),
        ("Checked at (UTC)", report["checked_at"]),
        ("Filters", filters),
        ("Rows in the standards list", report["database_rows"]),
        ("Rows in this export", report["exported_rows"]),
        ("Rows in the whole database", report["database_total"]),
    ]
    for status, count in sorted(report["by_status"]["database"].items()):
        exported = report["by_status"]["export"].get(status, 0)
        lines.append((f"{status}: list / export", f"{count} / {exported}"))
    lines += [
        ("Missing from export", ", ".join(report["missing_from_export"]) or "none"),
        ("Unexpected in export", ", ".join(report["unexpected_in_export"]) or "none"),
        ("Field mismatches", report["field_mismatch_count"]),
    ]
    sheet.append(["Check", "Value"])
    for column in (1, 2):
        header = sheet.cell(row=1, column=column)
        header.fill = HEADER_FILL
        header.font = HEADER_FONT
    for label, value in lines:
        _append(sheet, [label, value])
    workbook.save(path)


def _write_standards(sheet, rows) -> None:
    sheet.title = "Standards"
    sheet.append([header for header, _, _ in COLUMNS])
    for index, (_, _, width) in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
        cell = sheet.cell(row=1, column=index)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 34

    for standard, source in rows:
        _append(sheet, [_value(standard, source, attribute) for _, attribute, _ in COLUMNS])

    if rows:
        reference = f"A1:{get_column_letter(len(COLUMNS))}{len(rows) + 1}"
        table = Table(displayName="StandardsDatabase", ref=reference)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        sheet.add_table(table)
    sheet.freeze_panes = "A2"


def _value(standard: Standard, source: Source, attribute: str):
    if attribute == "_body":
        return source.body
    if attribute == "_created":
        return standard.created_at.strftime("%Y-%m-%d %H:%M") if standard.created_at else None
    if attribute == "_updated":
        return standard.updated_at.strftime("%Y-%m-%d %H:%M") if standard.updated_at else None
    return getattr(standard, attribute, None)


def _write_sources(sheet, db: Session) -> None:
    headers = ["Standards Body", "TC / SC Code", "Group Name", "Standard list (Source URL)",
               "Publish date policy", "Transport", "Active", "Last run"]
    sheet.append(headers)
    for index, width in enumerate((16, 24, 46, 80, 22, 12, 10, 20), start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
        cell = sheet.cell(row=1, column=index)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    for source in db.scalars(select(Source).order_by(Source.body, Source.tc_code)):
        _append(
            sheet,
            [
                source.body,
                source.tc_code,
                source.title,
                source.url,
                source.publish_date_policy,
                source.fetch_strategy,
                "Yes" if source.is_active else "No",
                source.last_run_at.strftime("%Y-%m-%d %H:%M") if source.last_run_at else "",
            ],
        )
    sheet.freeze_panes = "A2"


def _write_change_log(sheet, db: Session) -> None:
    headers = ["Run", "Date (UTC)", "Change", "Standard ID + Version", "Standards Body",
               "TC / SC Code", "English Title", "Old status", "New status",
               "Publication Type", "Relation", "Related standards"]
    sheet.append(headers)
    for index, width in enumerate((8, 20, 14, 34, 14, 20, 70, 14, 14, 22, 22, 40), start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
        cell = sheet.cell(row=1, column=index)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    query = (
        select(ChangeRecord, Source.body)
        .join(Source, Source.id == ChangeRecord.source_id)
        .order_by(ChangeRecord.id.desc())
    )
    for change, body in db.execute(query):
        details = change.details or {}
        related = (
            details.get("supersededBy")
            or details.get("supersedes")
            or ([details["amends"]] if details.get("amends") else [])
            or details.get("relatedStandards")
            or details.get("fields")
            or []
        )
        _append(
            sheet,
            [
                change.run_id,
                change.created_at.strftime("%Y-%m-%d %H:%M") if change.created_at else "",
                change.change_type,
                change.standard_id,
                body,
                change.tc_code,
                change.title_en,
                change.old_status,
                change.new_status,
                details.get("publicationType"),
                details.get("relation"),
                ", ".join(str(r) for r in related),
            ],
        )
    sheet.freeze_panes = "A2"
