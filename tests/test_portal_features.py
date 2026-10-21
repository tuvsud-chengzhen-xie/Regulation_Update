"""Portal features from test feedback: export verification, schedule, sources,
fetch-failure log and background runs. Uses the shared app database."""

from __future__ import annotations

import io
import time

import openpyxl

from app.services import ingest

ADMIN = {"X-API-Key": "test-admin-key"}


def ensure_data(client):
    if client.get("/api/standards?limit=1", headers=ADMIN).json()["total"] == 0:
        client.post("/api/runs", json={}, headers=ADMIN).raise_for_status()


def test_filtered_export_matches_the_standards_list_exactly(client):
    ensure_data(client)
    filters = "status=Current&body=CEN"
    listed = client.get(f"/api/standards?limit=1000&{filters}", headers=ADMIN).json()

    response = client.get(f"/api/export/standards.xlsx?{filters}", headers=ADMIN)
    assert response.status_code == 200
    assert response.headers["x-export-verified"] == "pass"
    assert int(response.headers["x-export-rows"]) == listed["total"]

    workbook = openpyxl.load_workbook(io.BytesIO(response.content))
    sheet = workbook["Standards"]
    exported = [row[0] for row in sheet.iter_rows(min_row=2, values_only=True)]
    assert exported == [item["standard_id"] for item in listed["items"]]
    check = {row[0]: row[1] for row in workbook["Export Check"].iter_rows(values_only=True)}
    assert check["Result"].startswith("PASS")


def test_export_keeps_scraped_text_as_text(client):
    ensure_data(client)
    item = client.get("/api/standards?limit=1&body=CEN", headers=ADMIN).json()["items"][0]
    client.patch(
        f"/api/standards/{item['id']}",
        json={"title_zh": "=HYPERLINK(\"http://evil\")\x07"},
        headers=ADMIN,
    ).raise_for_status()

    report = client.get(
        f"/api/export/verify?q={item['standard_id']}", headers=ADMIN
    ).json()
    assert report["ok"], report
    response = client.get(f"/api/export/standards.xlsx?q={item['standard_id']}", headers=ADMIN)
    sheet = openpyxl.load_workbook(io.BytesIO(response.content))["Standards"]
    row = next(r for r in sheet.iter_rows(min_row=2) if r[0].value == item["standard_id"])
    cell = row[2]
    assert cell.data_type == "s"
    assert cell.value == "=HYPERLINK(\"http://evil\")"


def test_schedule_is_configurable_from_the_portal(client):
    saved = client.put(
        "/api/runs/schedule",
        json={"enabled": True, "frequency": "weekly", "day_of_week": "wed", "hour": 3,
              "minute": 30, "start_date": "2099-01-01", "end_date": "2099-03-31"},
        headers=ADMIN,
    )
    assert saved.status_code == 200, saved.text
    body = saved.json()
    assert body["description"].startswith("Every Wed at 03:30 UTC")
    assert body["next_run_time"].startswith("2099-01-")
    assert client.get("/api/runs/schedule", headers=ADMIN).json()["frequency"] == "weekly"

    bad = client.put("/api/runs/schedule", json={"frequency": "cron", "cron": "not a cron"},
                     headers=ADMIN)
    assert bad.status_code == 422

    client.put("/api/runs/schedule", json={"enabled": False}, headers=ADMIN).raise_for_status()
    assert client.get("/api/runs/schedule", headers=ADMIN).json()["next_run_time"] is None


def test_coverage_reports_examples_against_targets(client):
    report = client.get("/api/runs/coverage", headers=ADMIN).json()
    assert set(report["observed"]) == {"added", "updated", "withdrawn"}
    assert set(report["by_type"]) >= {"added", "updated", "withdrawn", "superseded"}
    assert isinstance(report["sufficient"], bool)


def test_sources_can_be_added_only_on_the_bodies_own_sites(client):
    rejected = client.post(
        "/api/sources",
        json={"body": "ASTM", "tc_code": "F15.99", "url": "http://169.254.169.254/latest"},
        headers=ADMIN,
    )
    assert rejected.status_code == 422

    url = ("https://www.astm.org/membership-participation/technical-committees/"
           "committee-f15/subcommittee-f15/jurisdiction-f1543")
    result = client.post(
        "/api/sources/import",
        json=[{"body": "astm", "tc_code": "F15.43", "title": "Soccer Goal Safety", "url": url},
              {"body": "ASTM", "tc_code": "F15.42", "url": url}],
        headers=ADMIN,
    ).json()
    assert [s["tc_code"] for s in result["created"]] == ["F15.43"]
    assert result["created"][0]["fetch_strategy"] == "http"
    assert result["skipped"] and not result["errors"]


def test_failed_sources_are_listed_with_cause_and_health(client, payloads):
    payloads.set("CEN/TC 93", None)
    run = client.post("/api/runs", json={}, headers=ADMIN).json()
    failures = client.get("/api/runs/source-failures?latest=true", headers=ADMIN).json()
    tc93 = next(f for f in failures if f["tc_code"] == "CEN/TC 93")
    assert tc93["run_id"] == run["id"]
    assert tc93["cause"] and tc93["cause_label"] and tc93["error"]

    sources = {s["tc_code"]: s for s in client.get("/api/sources", headers=ADMIN).json()}
    assert sources["CEN/TC 93"]["last_status"] == "failed"
    assert sources["CEN/TC 93"]["consecutive_failures"] >= 1
    assert sources["CEN/TC 207"]["last_status"] == "success"


def test_runs_can_go_to_the_background_and_never_overlap(client):
    ingest._run_lock.acquire()
    try:
        busy = client.post("/api/runs", json={"background": True}, headers=ADMIN)
        assert busy.status_code == 409
    finally:
        ingest._run_lock.release()

    started = client.post("/api/runs", json={"background": True, "send_alerts": False},
                          headers=ADMIN)
    assert started.status_code == 202
    run_id = started.json()["id"]
    deadline = time.monotonic() + 60
    while time.monotonic() < deadline:
        run = client.get(f"/api/runs/{run_id}", headers=ADMIN).json()
        if run["status"] != "running":
            break
        time.sleep(0.2)
    assert run["status"] in {"success", "partial"}
    assert not client.get("/api/runs/active", headers=ADMIN).json()["busy"]


def test_standard_history_links_a_record_to_its_family(client):
    ensure_data(client)
    page = client.get("/api/standards?q=EN 16121:2023&limit=5", headers=ADMIN).json()
    corrigendum = next(i for i in page["items"] if i["standard_id"] == "EN 16121:2023/AC:2025")
    assert corrigendum["publication_type"] == "Corrigendum"
    history = client.get(f"/api/standards/{corrigendum['id']}/history", headers=ADMIN).json()
    assert {s["standard_id"] for s in history["family"]} >= {"EN 16121:2023",
                                                             "EN 16121:2023/AC:2025"}
    assert history["changes"] and history["changes"][0]["change_type"] == "added"


def source_id(client, tc_code):
    return next(s["id"] for s in client.get("/api/sources", headers=ADMIN).json()
                if s["tc_code"] == tc_code)


def test_a_blocked_site_can_be_updated_from_a_saved_page(client, payloads):
    from app.extractors.base import SAMPLE_DIR

    wstda = source_id(client, "WSTDA")
    client.patch(f"/api/sources/{wstda}", json={"fetch_strategy": "manual"},
                 headers=ADMIN).raise_for_status()
    payloads.set("WSTDA", None)  # the live fetch would fail; manual means it is not tried

    run = client.post("/api/runs", json={}, headers=ADMIN).json()
    result = next(r for r in run["results"] if r["source_id"] == wstda)
    assert result["status"] == "skipped"

    page = (SAMPLE_DIR / "wstda_products.html").read_bytes()
    uploaded = client.post(
        f"/api/sources/{wstda}/upload",
        files={"file": ("products.html", page, "text/html")},
        headers=ADMIN,
    )
    assert uploaded.status_code == 200, uploaded.text
    body = uploaded.json()
    assert body["trigger"] == "upload"
    assert body["results"][0]["status"] == "success"
    assert body["results"][0]["extracted_count"] > 0


def test_an_uploaded_verification_page_is_rejected(client):
    wstda = source_id(client, "WSTDA")
    page = b"<html><head><title>Just a moment...</title></head><body></body></html>"
    run = client.post(
        f"/api/sources/{wstda}/upload",
        files={"file": ("products.html", page, "text/html")},
        headers=ADMIN,
    ).json()
    result = run["results"][0]
    assert (result["status"], result["error_cause"]) == ("failed", "challenge")
    assert "verification page" in result["error"]


def test_iso_entries_still_under_publication_are_not_captured():
    from types import SimpleNamespace

    from app.extractors.iso import ISOExtractor

    html = """<table>
      <tr><td><div class="entry-name">ISO 4211-6:2025</div>
              <div class="entry-description">Furniture - Part 6</div></td><td>60.60</td></tr>
      <tr><td><div class="entry-name">ISO 4211-7</div> [Under development]
              <div class="entry-description">Furniture - Part 7</div></td><td>60.00</td></tr>
    </table> catalogue"""
    extractor = ISOExtractor(SimpleNamespace(tc_code="ISO/TC 136", url="https://x", options={}))
    result = extractor.parse(html.encode())
    assert [r.standard_id for r in result.rows] == ["ISO 4211-6:2025"]
    assert not result.issues
