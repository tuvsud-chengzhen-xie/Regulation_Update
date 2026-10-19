"""End-to-end API behaviour: access control, monitoring, alerts, export."""

from __future__ import annotations

import io

import openpyxl

ADMIN = {"X-API-Key": "test-admin-key"}


def test_health_is_public(client):
    response = client.get("/health")
    assert response.status_code == 200
    assert response.json()["status"] == "ok"


def test_api_requires_a_key(client):
    assert client.get("/api/standards").status_code == 401
    assert client.get("/api/standards", headers={"X-API-Key": "nope"}).status_code == 401


def test_bootstrap_seeded_all_sources(client):
    sources = client.get("/api/sources", headers=ADMIN).json()
    bodies = {s["body"] for s in sources}
    assert bodies == {"CEN", "ISO", "ASTM", "ASME", "BIFMA", "ALI", "WSTDA"}


def test_first_run_captures_standards_as_current(client):
    run = client.post("/api/runs", json={}, headers=ADMIN).json()
    assert run["status"] in {"success", "partial"}
    assert run["added_count"] > 100

    page = client.get("/api/standards?limit=1", headers=ADMIN).json()
    assert page["total"] == run["added_count"]
    assert page["items"][0]["status"] == "Current"

    stats = client.get("/api/standards/statistics", headers=ADMIN).json()
    assert stats["by_status"]["Current"] == page["total"]


def test_a_source_that_cannot_be_fetched_is_reported_not_hidden(client, payloads):
    """An unreachable source is a gap in the database, and must be visible."""
    payloads.set("CEN/TC 52", None)
    run = client.post("/api/runs", json={}, headers=ADMIN).json()

    failed = [r for r in run["results"] if r["status"] == "failed"]
    assert failed and all(r["error"] for r in failed)

    issues = client.get(f"/api/runs/{run['id']}/issues?scope=source", headers=ADMIN).json()
    assert any(i["tc_code"] == "CEN/TC 52" for i in issues)
    # Nothing was withdrawn on the strength of a failed fetch.
    assert run["withdrawn_count"] == 0


def test_second_run_is_a_no_op(client):
    run = client.post("/api/runs", json={}, headers=ADMIN).json()
    assert run["added_count"] == 0
    assert run["withdrawn_count"] == 0


def test_non_admin_cannot_trigger_runs_or_export(client):
    created = client.post(
        "/api/users",
        json={"email": "engineer@example.com", "full_name": "Engineer", "role": "user"},
        headers=ADMIN,
    ).json()
    user_headers = {"X-API-Key": created["api_key"]}

    assert client.post("/api/runs", json={}, headers=user_headers).status_code == 403
    assert client.get("/api/export/standards.xlsx", headers=user_headers).status_code == 403
    # ...but they can read the database and manage their own subscriptions.
    assert client.get("/api/standards?limit=1", headers=user_headers).status_code == 200


def test_subscription_alerting_end_to_end(client, payloads):
    created = client.post(
        "/api/users",
        json={"email": "furniture@example.com", "full_name": "Furniture expert", "role": "user"},
        headers=ADMIN,
    ).json()
    user_headers = {"X-API-Key": created["api_key"]}

    client.post(
        "/api/subscriptions",
        json={"type": "committee", "value": "CEN/TC 207"},
        headers=user_headers,
    ).raise_for_status()

    # Force a change: serve CEN/TC 207 a different catalogue so its standards
    # disappear, then restore it so the next run re-adds them.
    payloads.set("CEN/TC 207", "cen_tc93.xml")

    run = client.post("/api/runs", json={}, headers=ADMIN).json()
    assert run["withdrawn_count"] > 0
    assert run["alerts_sent"] >= 1

    alerts = client.get("/api/subscriptions/alerts", headers=user_headers).json()
    assert alerts
    body = client.get(f"/api/subscriptions/alerts/{alerts[0]['id']}", headers=user_headers).json()
    assert "CEN/TC 207" in body["body"]
    assert "matched your committee subscription" in body["body"]

    payloads.set("CEN/TC 207", "cen_tc207.xml")
    restored = client.post("/api/runs", json={}, headers=ADMIN).json()
    assert restored["reinstated_count"] > 0


def test_change_log_records_both_directions(client):
    changes = client.get("/api/runs/changes?limit=500", headers=ADMIN).json()
    types = {c["change_type"] for c in changes}
    assert {"added", "withdrawn", "reinstated"} <= types


def test_excel_export_uses_the_database_format_columns(client):
    response = client.get("/api/export/standards.xlsx", headers=ADMIN)
    assert response.status_code == 200

    workbook = openpyxl.load_workbook(io.BytesIO(response.content))
    sheet = workbook["Standards"]
    headers = [cell.value for cell in sheet[1]]
    assert headers[:7] == [
        "标准号+年份\nStandard ID + Version",
        "归口单位代码\nTC / SC Code",
        "标准中文名称\nStandard Chinese Title",
        "标准英文名称\nStandard English Title",
        "发布日期\nPublish Date",
        "实施日期\nImplement Date",
        "标准状态\nStatus",
    ]
    assert sheet.max_row > 100
    assert {"Standards", "Sources", "Change Log"} <= set(workbook.sheetnames)


def test_admin_can_curate_a_chinese_title(client):
    page = client.get("/api/standards?limit=1&status=Current", headers=ADMIN).json()
    standard_id = page["items"][0]["id"]
    response = client.patch(
        f"/api/standards/{standard_id}", json={"title_zh": "家具 — 座椅"}, headers=ADMIN
    )
    assert response.status_code == 200
    assert response.json()["title_zh"] == "家具 — 座椅"


def test_design_system_and_app_shell_are_served(client):
    """The UI is only usable if the CPS CoE design system mounts correctly."""
    index = client.get("/")
    assert index.status_code == 200
    body = index.text
    # Shell markup from the design system's skeleton.
    assert 'class="topbar"' in body
    assert "by CPS Center of Excellence" in body
    assert '<link rel="stylesheet" href="/ds/tokens.css">' in body

    for asset in ("/ds/tokens.css", "/ds/shell.css", "/ds/shell.js",
                  "/ds/nudges.css", "/ds/nudges.js",
                  "/static/app.css", "/static/app.js", "/static/logo.png"):
        assert client.get(asset).status_code == 200, asset

    # The AI identity must NOT be loaded: this app has no Layer-3 AI tool.
    assert "ai-overlay" not in body
    assert "AIOverlay" not in body


def test_statistics_feed_the_dashboard_kpis(client):
    stats = client.get("/api/standards/statistics", headers=ADMIN).json()
    assert {"total", "by_status", "by_body", "sources", "changes_total"} <= set(stats)
    assert stats["sources"]["active"] <= stats["sources"]["total"]
    assert stats["changes_total"] >= 0


def test_issue_endpoints_report_what_was_not_captured(client, payloads):
    """The Issues screen is driven entirely by these three endpoints."""
    payloads.set("CEN/TC 93", None)          # a whole source goes missing
    run = client.post("/api/runs", json={}, headers=ADMIN).json()

    summary = client.get("/api/runs/issues/summary?latest=true", headers=ADMIN).json()
    assert summary["run_id"] == run["id"]
    assert summary["total"] == summary["lost_records"] + summary["incomplete_records"] \
        + summary["failed_sources"]
    assert summary["failed_sources"] >= 1
    assert summary["by_scope"] and summary["by_reason"] and summary["by_body"]

    issues = client.get("/api/runs/issues?latest=true", headers=ADMIN).json()
    assert len(issues) == summary["total"]
    for issue in issues:
        assert issue["scope"] in {"record", "field", "source"}
        assert issue["detail"] and issue["body"] and issue["tc_code"]

    lost = client.get("/api/runs/issues?latest=true&scope=record", headers=ADMIN).json()
    assert all(i["scope"] == "record" for i in lost)
    assert len(lost) == summary["lost_records"]

    per_run = client.get(f"/api/runs/{run['id']}/issues", headers=ADMIN).json()
    assert len(per_run) == summary["total"]
    assert client.get("/api/runs/999999/issues", headers=ADMIN).status_code == 404


def test_sources_expose_their_transport(client):
    sources = client.get("/api/sources", headers=ADMIN).json()
    by_body = {s["body"]: s["fetch_strategy"] for s in sources}
    # The four sites that cannot be read with a plain HTTP client.
    for body in ("ISO", "BIFMA", "ALI", "WSTDA"):
        assert by_body[body] == "browser", body
    for body in ("CEN", "ASTM", "ASME"):
        assert by_body[body] == "http", body
